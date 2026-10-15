"""Coverage, offline routing and local source boundaries; real inference replay is separate."""
import asyncio
from io import BytesIO
import json
from pathlib import Path
import tempfile
from unittest.mock import patch
import pytest
from PIL import Image
from federation.core import FederationStore
from federation.tool_files import put_file,get_file,search_files
from federation.document_tools import read_document,parsed_document
from federation.tool_services import save_service,status,model_config,record_execution
from federation.workflow_tools import execute
from federation.local_files import import_local

@pytest.fixture
def store(tmp_path):
    s=FederationStore(tmp_path/'data');s.init();return s

def read(s,f,question='说明全文',start='0',limit='0'):
    return asyncio.run(read_document(s,'a'*32,f['id'],question,start,limit))

def test_full_document_tail_ranges_cache_and_hash(store):
    text='资料正文\n'*33000+'FINAL-78291'
    f=put_file(store,'a'*32,'long.txt',text.encode())
    r=read(store,f,start=str(len(text)-50),limit='200')
    assert 'FINAL-78291' in r['text'] and r['source']['next_start'] is None
    assert r['source']['total_chars']>160000 and not r['source']['extraction_truncated']
    with pytest.raises(ValueError):read(store,f,start=str(len(text)+100),limit='0')
    _,p=get_file(store,'a'*32,f['id']);p.write_bytes(b'changed')
    with pytest.raises(ValueError):read(store,f)

def test_long_document_maps_every_range_without_cloud(store):
    text='first '+('x'*36000)+' TAIL-9281';f=put_file(store,'a'*32,'long.txt',text.encode());calls=[]
    async def summarize(question,images=(),**kw):
        calls.append(question);return {'text':'尾部：TAIL-9281' if 'TAIL-9281' in question else '中间正文','model':'local-fixture','usage':{'input_tokens':10,'output_tokens':5}}
    with patch('federation.local_models.local_reply',side_effect=summarize):r=read(store,f)
    assert len(calls)==4 and r['source']['end']==r['source']['total_chars'] and r['source']['mode']=='local_chunk_summary'
    assert 'TAIL-9281' in r['sections'][-1]['text'] and not r['source']['truncated']
    with patch('federation.local_models.local_reply',side_effect=AssertionError('Cached chunks must not invoke a model again')):cached=read(store,f)
    assert cached['sections']==r['sections'] and all(u['calls']==0 for u in cached['model_usage'])

def test_medium_document_raw_ranges_fit_context_and_reach_tail(store):
    f=put_file(store,'a'*32,'medium.txt',('中'*12000+'TAIL-9281').encode())
    parts=[];start='0'
    for _ in range(4):
        r=read(store,f,start=start,limit='18000')
        assert len(r['text'])<=6000
        parts.append(r['text'])
        if r['source']['next_start'] is None:break
        start=str(r['source']['next_start'])
    assert 'TAIL-9281' in ''.join(parts) and not r['source']['truncated']

def test_word_pptx_excel_order_notes_and_formulas(store):
    from docx import Document
    doc=Document();doc.add_paragraph('before');doc.add_table(rows=1,cols=1).cell(0,0).text='TABLE-288';doc.add_paragraph('TAIL-9281');doc.sections[0].header.paragraphs[0].text='HEADER-728'
    b=BytesIO();doc.save(b);f=put_file(store,'a'*32,'file.docx',b.getvalue());r=read(store,f)
    assert r['text'].index('before')<r['text'].index('TABLE-288')<r['text'].index('TAIL-9281') and 'HEADER-728' in r['text']
    from pptx import Presentation
    deck=Presentation()
    for value in ['FIRST-728','LAST-9281']:
        slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text=value
    slide.notes_slide.notes_text_frame.text='NOTE-288';b=BytesIO();deck.save(b)
    f=put_file(store,'a'*32,'file.pptx',b.getvalue());r=read(store,f)
    assert all(t in r['text'] for t in ['FIRST-728','LAST-9281','NOTE-288']) and len(r['source']['units'])==2
    assert search_files(store,'a'*32,'9281',f['id'])['status']=='completed'
    from openpyxl import Workbook
    book=Workbook();book.active.append(['amount',288,'=B1*2']);b=BytesIO();book.save(b)
    f=put_file(store,'a'*32,'file.xlsx',b.getvalue());r=read(store,f);assert '=B1*2' in r['text'] and '288' in r['text']

def test_scan_coverage_and_offline_local_vision(store):
    from pypdf import PdfWriter
    writer=PdfWriter();writer.add_blank_page(100,100);writer.add_blank_page(100,100);b=BytesIO();writer.write(b)
    f=put_file(store,'a'*32,'scan.pdf',b.getvalue());r=read(store,f)
    assert r['status']=='no_results' and r['source']['empty_units']==['第1页','第2页']
    save_service(store,'vision',{'provider':'local-ollama','enabled':True})
    assert model_config(store,'vision')['provider']=='local-ollama'
    calls=[]
    async def vision(question,images=(),**kw):calls.append(images);return {'text':'实际页面','model':'local-fixture','usage':{}}
    context={'base':str(store.base),'room_id':'a'*32,'tools_enabled':True,'network_allowed':False}
    with patch('federation.local_models.local_reply',side_effect=vision):
        r=asyncio.run(execute(context,'image_analyze',{'file_id':f['id'],'question':'说明','start':'2','count':'1'}))
    assert len(calls)==1 and r['source']['start_page']==2 and r['source']['end_page']==2
    save_service(store,'vision',{'enabled':False})
    r=asyncio.run(execute(context,'image_analyze',{'file_id':f['id'],'question':'说明','start':'1','count':'1'}))
    assert r['status']=='unavailable'

def test_local_import_read_only_symlinks_private_paths_and_scope(store):
    with tempfile.TemporaryDirectory(dir=Path.home(),prefix='openwiki-import-test-') as folder:
        root=Path(folder);p=root/'document.txt';p.write_text('KEEP-288');(root/'other.txt').write_text('another');(root/'.env.secret.txt').write_text('test-only')
        (root/'link.txt').symlink_to(p)
        r=import_local(store,'a'*32,str(root));assert len(r['files'])==2 and r['skipped_count']==2 and p.read_text()=='KEEP-288'
        with pytest.raises(ValueError):import_local(store,'a'*32,str(root/'link.txt'))
        with pytest.raises(ValueError):import_local(store,'a'*32,str(root/'.env.secret.txt'))
        with pytest.raises(FileNotFoundError):get_file(store,'b'*32,r['files'][0]['id'])
    with pytest.raises(ValueError):import_local(store,'a'*32,'/etc/passwd')

def test_media_types_limits_disabled_tools_and_fingerprint(store):
    assert put_file(store,'a'*32,'audio.mp3',b'fixture')['kind']=='audio'
    f=put_file(store,'a'*32,'movie.mp4',b'fixture');assert f['kind']=='video'
    with pytest.raises(ValueError):put_file(store,'a'*32,'audio.wav',b'x'*50000001)
    save_service(store,'video',{'enabled':False})
    r=asyncio.run(execute({'base':str(store.base),'room_id':'a'*32,'tools_enabled':True},'video_analyze',{'file_id':f['id'],'question':'内容','start':'0','duration':'300'}));assert r['status']=='unavailable'
    record_execution(store,'files',{'status':'completed'});assert status(store)['files']['verified']
    assert 'api_key' not in json.dumps(status(store))

def test_attachment_tasks_do_not_route_to_cloud_tool_model(store):
    from federation.agent_templates import initialize_standards
    from federation.app import isolated_ask_environment
    from federation.report_model import _write_private
    initialize_standards(store)
    _write_private(store.base/'private/report-model.json',{'provider':'openai-compatible','model':'fixture','base_url':'https://example.com/v1','api_key':'test-only','responses_api':True})
    f=put_file(store,'a'*32,'file.txt',b'text')
    job={'id':'c'*32,'space':'central','question':'分析文件','room_id':'a'*32,'delegation_phase':'consultation','tool_request':True,'files':[f]}
    env=isolated_ask_environment(store,job)
    assert not job['tool_policy']['use_tool_model'] and 'FEDERATION_TOOL_AGENT_MODEL_CONFIG' not in env

def test_media_time_windows_reject_nonfinite_and_overlong():
    from federation.audio_video import window
    for start,duration in [('nan','10'),('0','inf'),('-1','1'),('0','301'),('500','10')]:
        with pytest.raises(ValueError):window(start,duration,100)
    assert window('80','30',100)==(80,100)

def test_media_preview_is_same_origin_with_inline_and_range(tmp_path):
    from fastapi.testclient import TestClient
    from federation.app import create_app
    app=create_app(base=tmp_path/'preview')
    with TestClient(app) as client:
        assert "media-src 'self'" in client.get('/').headers['content-security-policy']
        room=client.post('/api/chat/rooms',json={'mode':'direct','participants':['central']}).json()
        file=put_file(app.state.store,room['id'],'speech.wav',b'fixture-media')
        response=client.get(file['url'],headers={'Range':'bytes=0-3'})
        assert response.status_code==206 and response.content==b'fixt'
        assert response.headers['content-disposition'].startswith('inline;')
        assert response.headers['content-type']=='audio/x-wav'
        assert response.headers['x-content-type-options']=='nosniff'
