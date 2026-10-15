"""Disposable bounded text parser, never executes document macros or embedded code."""
from io import BytesIO
import resource
import sys
resource.setrlimit(resource.RLIMIT_CPU,(15,15))
if sys.platform!='darwin':resource.setrlimit(resource.RLIMIT_AS,(512*1024*1024,512*1024*1024))
try:
    raw=sys.stdin.buffer.read(5_000_001)
    if len(raw)>5_000_000:raise ValueError('文件超过5MB')
    if sys.argv[1]=='.pdf':
        from pypdf import PdfReader
        doc=PdfReader(BytesIO(raw))
        if doc.is_encrypted:raise ValueError('PDF已加密，请先解密')
        if len(doc.pages)>100:raise ValueError('PDF超过100页，请先拆分')
        blocks=[]
        for index,page in enumerate(doc.pages):
            content=page.get_contents()
            if content is not None and len(content.get_data())>4_000_000:raise ValueError('PDF页内容过大，请拆分')
            blocks.append(page.extract_text() or '')
        if not any(block.strip() for block in blocks):raise ValueError('扫描件未读到文字，需先OCR')
        blocks=['## 第'+str(index+1)+'页\n\n'+text for index,text in enumerate(blocks)]
        text='\n\n'.join(blocks)
    elif sys.argv[1]=='.docx':
        from zipfile import ZipFile
        with ZipFile(BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist())>20_000_000:raise ValueError('Word解压大小超过限制')
        from docx import Document
        doc=Document(BytesIO(raw));text='\n\n'.join(p.text for p in doc.paragraphs)
        for table in doc.tables:text+='\n\n'+'\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows)
    else:raise ValueError('文档格式无效')
    if len(text.encode())>80000:raise ValueError('提取正文超过80000字节，请拆分')
    print(text,end='')
except Exception as exc:
    print(str(exc)[:200]);sys.exit(1)
