"""Downloadable research artifacts, isolated from Wiki and published knowledge."""
import hashlib
import json
import re
from pathlib import Path

from .core import _now, _write_json


def _run_id(value):
    if not isinstance(value, str) or not re.fullmatch(r'[a-f0-9]{32}', value):
        raise ValueError('报告标识无效')
    return value


def _plain(text):
    text = re.sub(r'\[([^\]]+)\]\((https?://[^\s)]+)\)', r'\1（\2）', text)
    return re.sub(r'\*\*|__|`', '', text)


def write_docx(markdown, destination):
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(22)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number']:
        style = doc.styles[name]
        style.font.name = 'Noto Sans CJK SC'
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
            fonts.attrib.pop(qn('w:' + key), None)
        for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            fonts.set(qn('w:' + key), 'Noto Sans CJK SC')
        for border in style._element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    doc.styles['Title'].font.size = Pt(23)
    for name, size in [('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 12)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].paragraph_format.space_before = Pt(14)
        doc.styles[name].paragraph_format.space_after = Pt(7)
    lines = markdown.splitlines()
    index = 0
    in_code = False
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line or re.fullmatch(r'[-*_]{3,}', line):
            continue
        if line.startswith('```'):
            in_code = not in_code
            continue
        if in_code:
            p = doc.add_paragraph(line)
            for run in p.runs:
                run.font.name = 'Courier New'
                run.font.size = Pt(9)
            continue
        heading = re.match(r'^(#{1,4})\s+(.+)', line)
        if heading:
            level = len(heading[1])
            doc.add_paragraph(_plain(heading[2]), style='Title' if level == 1 and not doc.paragraphs else 'Heading ' + str(min(3, max(1, level - 1))))
            continue
        if line.startswith('|') and index < len(lines) and re.match(r'^\s*\|?\s*:?-{3}', lines[index]):
            headers = [cell.strip() for cell in line.strip('|').split('|')]
            separator = [cell.strip() for cell in lines[index].strip().strip('|').split('|')]
            index += 1
            rows = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                rows.append([cell.strip() for cell in lines[index].strip().strip('|').split('|')])
                index += 1
            valid_separator = len(separator) == len(headers) and all(re.fullmatch(r':?-{3,}:?', cell) for cell in separator)
            if not valid_separator or any(len(values) != len(headers) for values in rows):
                # Preserve malformed model Markdown instead of silently losing
                # cells or presenting text under an incorrect column heading.
                for values in [headers] + rows:
                    doc.add_paragraph(_plain(' | '.join(values)))
                continue
            table = doc.add_table(rows=1, cols=len(headers))
            table.autofit = True
            for row_number, values in enumerate([headers] + rows):
                cells = table.rows[0].cells if row_number == 0 else table.add_row().cells
                for column, cell in enumerate(cells):
                    cell.text = _plain(values[column]) if column < len(values) else ''
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    props = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement('w:tcBorders')
                    for side in ['top', 'left', 'bottom', 'right']:
                        edge = OxmlElement('w:' + side)
                        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                            edge.set(qn('w:' + key), value)
                        borders.append(edge)
                    props.append(borders)
                    margins = OxmlElement('w:tcMar')
                    for side in ['top', 'left', 'bottom', 'right']:
                        margin = OxmlElement('w:' + side)
                        margin.set(qn('w:w'), '100'); margin.set(qn('w:type'), 'dxa'); margins.append(margin)
                    props.append(margins)
                    if row_number == 0:
                        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E9EDF1'); props.append(shade)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(3)
                        for run in paragraph.runs:
                            run.font.size = Pt(10)
                            run.bold = row_number == 0
                if row_number == 0:
                    repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        bullet = re.match(r'^[-*+]\s+(.+)', line)
        numbered = re.match(r'^\d+[.)、]\s*(.+)', line)
        doc.add_paragraph(_plain(bullet[1] if bullet else numbered[1] if numbered else line.lstrip('> ')),
                          style='List Bullet' if bullet else 'List Number' if numbered else None)
    doc.core_properties.title = next((p.text for p in doc.paragraphs if p.text), '产品分析报告')[:120]
    doc.core_properties.author = '知识库'
    doc.save(str(destination))


def save_report(store, run_id, content, research=None, generation=None):
    _run_id(run_id)
    if not isinstance(content, str) or not content.strip() or len(content) > 60000 or '\0' in content:
        raise ValueError('报告正文无效')
    root = store._safe(store.base, 'reports/' + run_id)
    root.mkdir(parents=True, exist_ok=False, mode=0o700)
    title = next((line.lstrip('# ').strip() for line in content.splitlines() if line.strip()), '产品分析报告')[:100]
    body = content.strip()
    if not body.startswith('# '):
        title = '产品分析报告'
        body = '# ' + title + '\n\n' + body
    sources = (research or {}).get('sources') or []
    if sources:
        body += '\n\n## 本轮来源记录\n\n'
        for source in sources:
            body += '- ' + str(source.get('title', '公开资料')) + ' ' + str(source.get('url', '')) + '\n'
        body += '\n读取时间 ' + str((research or {}).get('finished_at', _now())) + '\n'
    artifacts = []
    for extension in ['md', 'docx']:
        file = root / ('report.' + extension)
        if extension == 'md':
            file.write_text(body + '\n', encoding='utf-8')
        else:
            write_docx(body, file)
        file.chmod(0o600)
        data = file.read_bytes()
        artifacts.append({'id': run_id, 'format': extension, 'title': title,
                          'filename': title + '.' + extension, 'bytes': len(data),
                          'sha256': hashlib.sha256(data).hexdigest(),
                          'url': '/api/reports/' + run_id + '/' + extension})
    _write_json(root / 'manifest.json', {'id': run_id, 'created_at': _now(), 'artifacts': artifacts,
                                       'generation':generation or {}})
    return artifacts


def report_file(store, run_id, extension):
    _run_id(run_id)
    if extension not in {'md', 'docx'}:
        raise ValueError('报告格式无效')
    root = store._safe(store.base, 'reports/' + run_id)
    manifest = json.loads(store._safe(root, 'manifest.json').read_text(encoding='utf-8'))
    artifact = next((item for item in manifest['artifacts'] if item['id'] == run_id and item['format'] == extension), None)
    file = store._safe(root, 'report.' + extension)
    if not artifact or not file.is_file() or hashlib.sha256(file.read_bytes()).hexdigest() != artifact['sha256']:
        raise ValueError('报告文件与登记版本不一致')
    return file, artifact
